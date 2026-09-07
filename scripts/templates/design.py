"""Design tokens and PPTX drawing helpers for the CubeOffice original templates.

Every template is described by a `Design`: a palette, a type pairing, a corner
radius and a cover motif. The slide recipes in `recipes.py` read nothing else,
so two templates that share a recipe still look unrelated. All artwork is drawn
from these tokens as plain OOXML shapes -- the templates embed no photographs,
no icon fonts and no third-party assets.
"""

from __future__ import annotations

from dataclasses import dataclass, field

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# 16:9 at the size PowerPoint itself writes for a widescreen deck.
SLIDE_WIDTH = Emu(12192000)
SLIDE_HEIGHT = Emu(6858000)
SLIDE_WIDTH_IN = 13.3333
SLIDE_HEIGHT_IN = 7.5

# One margin for every recipe, so slides from different templates still stack
# into a coherent deck when a user mixes them.
MARGIN = 0.92
CONTENT_WIDTH = SLIDE_WIDTH_IN - 2 * MARGIN


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.lstrip("#").upper())


def mix(a: str, b: str, weight: float) -> str:
    """Blend two hex colors. `weight` is how much of `b` to take."""
    a_c = a.lstrip("#")
    b_c = b.lstrip("#")
    out = []
    for i in (0, 2, 4):
        channel = round(
            int(a_c[i : i + 2], 16) * (1 - weight) + int(b_c[i : i + 2], 16) * weight
        )
        out.append(max(0, min(255, channel)))
    return "%02X%02X%02X" % tuple(out)


@dataclass(frozen=True)
class Palette:
    """Named roles rather than raw colors, so recipes never pick a swatch."""

    canvas: str  # slide background
    ink: str  # primary text
    muted: str  # secondary text
    accent: str  # the one loud color
    accent_ink: str  # text that sits on `accent`
    surface: str  # cards and panels on `canvas`
    rule: str  # hairlines and dividers
    cover_canvas: str  # cover / section background
    cover_ink: str  # text on `cover_canvas`

    def tint(self, role: str, weight: float) -> str:
        """A role blended toward the canvas -- for washes and quiet fills."""
        return mix(getattr(self, role), self.canvas, weight)


@dataclass(frozen=True)
class TypeScale:
    heading_font: str
    body_font: str
    display: float  # cover title
    title: float  # slide title
    section: float  # section-divider number/label
    lead: float  # intro paragraph
    body: float
    caption: float
    stat: float
    heading_bold: bool = True
    display_caps: bool = False
    title_caps: bool = False
    heading_spacing: float = 0.0  # extra character spacing, points


@dataclass(frozen=True)
class Design:
    key: str
    name_en: str
    name_zh: str
    tagline_en: str
    tagline_zh: str
    category_en: str
    category_zh: str
    palette: Palette
    type_scale: TypeScale
    motif: str
    radius: float = 0.0  # card corner radius, inches; 0 draws square corners
    content: dict = field(default_factory=dict)


# --------------------------------------------------------------------------
# Shape helpers. Each returns the created shape so a recipe can keep tuning it.
# --------------------------------------------------------------------------


def _no_line(shape) -> None:
    shape.line.fill.background()


def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0, radius=0.0):
    shape_kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        # `adjustments` is a fraction of the shorter side, not an absolute size.
        shape.adjustments[0] = min(0.5, radius / min(w, h))
    shape.shadow.inherit = False
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_w)
    else:
        _no_line(shape)
    shape.text_frame.text = ""
    return shape


def ellipse(slide, x, y, w, h, fill=None, line=None, line_w=1.0):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.shadow.inherit = False
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_w)
    else:
        _no_line(shape)
    return shape


def triangle(slide, x, y, w, h, fill, rotation=0.0):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_TRIANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.shadow.inherit = False
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    _no_line(shape)
    shape.rotation = rotation
    return shape


def parallelogram(slide, x, y, w, h, fill, skew=0.28):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.PARALLELOGRAM, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.shadow.inherit = False
    shape.adjustments[0] = skew
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    _no_line(shape)
    return shape


def rule(slide, x, y, w, color, weight=1.0):
    """A hairline drawn as a filled rectangle so its weight survives scaling."""
    return rect(slide, x, y, w, weight / 72.0, fill=color)


def vrule(slide, x, y, h, color, weight=1.0):
    return rect(slide, x, y, weight / 72.0, h, fill=color)


def textbox(
    slide,
    x,
    y,
    w,
    h,
    *,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    wrap=True,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = wrap
    frame.vertical_anchor = anchor
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.paragraphs[0].alignment = align
    return box


def write(
    frame,
    text,
    *,
    font,
    size,
    color,
    bold=False,
    italic=False,
    caps=False,
    spacing=0.0,
    line_spacing=1.15,
    space_before=0.0,
    space_after=0.0,
    align=None,
    first=False,
):
    """Append a paragraph. `first` reuses the empty paragraph a frame starts with."""
    paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
    paragraph.line_spacing = line_spacing
    if space_before:
        paragraph.space_before = Pt(space_before)
    if space_after:
        paragraph.space_after = Pt(space_after)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text.upper() if caps else text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    if spacing:
        # python-pptx has no spacing setter; `spc` is in 1/100 pt.
        run.font._rPr.set("spc", str(int(round(spacing * 100))))
    return paragraph


def bullets(frame, items, *, font, size, color, bullet_color, line_spacing=1.3, gap=9.0):
    """A bulleted list drawn with a leading glyph.

    A real `buChar` list would inherit its indent from the layout, which these
    templates deliberately do not define, so the marker is part of the run.
    """
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.line_spacing = line_spacing
        if index:
            paragraph.space_before = Pt(gap)
        marker = paragraph.add_run()
        marker.text = "—  "
        marker.font.name = font
        marker.font.size = Pt(size)
        marker.font.color.rgb = rgb(bullet_color)
        marker.font.bold = True
        body = paragraph.add_run()
        body.text = item
        body.font.name = font
        body.font.size = Pt(size)
        body.font.color.rgb = rgb(color)


def background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)
